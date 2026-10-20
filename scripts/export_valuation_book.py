"""Bảng định giá chuẩn VN100 → Excel nhiều sheet (chỉ đọc DB, không gọi AI).

Sheet: Hướng dẫn · Tổng hợp · Giá mục tiêu CTCK · Giả định · Số liệu TC ·
Phương pháp · Chú thích cờ · Chất lượng dữ liệu.
Giá trị hợp lý = đường định giá DUY NHẤT (batch.value_ticker, cùng lõi với Streamlit).
Upside, khuyến nghị, lệch CTCK, bội số... là CÔNG THỨC Excel để người dùng sửa đầu vào.

  python -m scripts.export_valuation_book [đường_dẫn_ra.xlsx]
Sau đó chạy recalc.py của skill xlsx để có giá trị cache (xem docs/AGENT_OPERATING_MANUAL.md).
"""

from __future__ import annotations

import csv
import datetime
import statistics
import sys
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from valuation.calibration.consensus_view import (
    DEFAULT_STALE_AFTER_DAYS,
    get_consensus_view,
)
from valuation.config import PROJECT_ROOT, load_defaults
from valuation.data_access.corporate_actions import _today_vn
from valuation.data_access.repo import build_company_data
from valuation.db.models import Consensus, FinancialsQuarterly, PricesDaily, Ticker
from valuation.db.session import SessionLocalRead
from valuation.engine.batch import value_ticker
from valuation.engine.decision_engine import REC_LABELS_VI
from valuation.engine.flag_descriptions import describe_flags
from valuation.engine.models.bank_general import converged_roe
from valuation.engine.sector_router import route
from valuation.engine.valuate import valuate
from valuation.ingest.universe import get_vn100_symbols
from valuation.output.gsheets_exporter import _METHOD_LABEL

WINDOW_DAYS = 180
CFG = load_defaults()
BAND = float((CFG.get("audit") or {}).get("consensus_review_band", 0.25))
TRIM_FLOOR = float((CFG.get("margin_of_safety") or {}).get("trim_floor", -0.10))
PRICE_MOVE_NOTE = (
    BAND  # giá lệch > band so với lúc CTCK viết → giá mục tiêu có thể lỗi thời
)
BANK_NI_DROP_NOTE = 0.8  # LN dự phóng năm 1 < 80% TTM → giải thích giả định

_BLOCK_REASON = {
    "VALUATION_PROXY": "Mô hình proxy (RNAV/SOTP chưa có dữ liệu dự án/mảng)",
    "NOT_RATED": "Mô hình tự từ chối định giá",
    "PROXY_IMPLAUSIBLE": "Proxy lệch phi lý so với thị giá",
    "UPSIDE_EXTREME_REVIEW": "Upside > +150% — chờ rà soát giả định",
    "DOWNSIDE_EXTREME_REVIEW": "Downside < −60% — chờ rà soát giả định",
    "AI_INPUTS_UNVERIFIED": "Đầu vào AI chưa được người duyệt",
    "NEGATIVE_EQUITY_VALUE_DCF": "Giá trị vốn CSH âm (DCF)",
    "NEGATIVE_EQUITY_VALUE_EV_EBITDA": "Giá trị vốn CSH âm (EV/EBITDA)",
}

ARIAL = "Arial"
F_BASE = Font(name=ARIAL, size=10)
F_INPUT = Font(name=ARIAL, size=10, color="0000FF")
F_LINK = Font(name=ARIAL, size=10, color="008000")
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=ARIAL, size=14, bold=True)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
WRAP = Alignment(wrap_text=True, vertical="top")
FMT_VND = '#,##0;(#,##0);"-"'
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_X = '0.0"x";(0.0"x");"-"'
FMT_DATE = "dd/mm/yyyy"


# ---------------------------------------------------------------- dữ liệu
def _repair_status(fixed: set[str]) -> dict[str, str]:
    """Trạng thái sửa item_id trùng (scripts/repair_duplicate_bs_items).

    `fixed` = mã đã có dòng BS tách hậu tố "__" trong DB (bằng chứng đã sửa). Log chỉ dùng
    để biết mã nào lần thử gần nhất bị lỗi mạng — log cũ từng bị ghi đè cùng tên file nên
    không dùng log để khẳng định "đã sửa".
    """
    status: dict[str, str] = {}
    for f in sorted(
        (PROJECT_ROOT / "reports" / "audit").glob("bs_duplicate_repair_*.csv"),
        key=lambda p: p.stat().st_mtime,
    ):
        with open(f, encoding="utf-8-sig") as fh:
            for row in csv.DictReader(fh):
                if (row.get("action") or "").startswith("FETCH_ERROR"):
                    status[row["ticker"]] = "Chưa sửa được (lỗi mạng nguồn)"
                else:
                    status.pop(row["ticker"], None)
    status.update({t: "Đã sửa (có dòng tách item_id trùng)" for t in fixed})
    return status


def _fixed_tickers(db) -> set[str]:
    rows = (
        db.query(FinancialsQuarterly.ticker)
        .filter(
            FinancialsQuarterly.statement == "BS",
            FinancialsQuarterly.line_item.like("%\\_\\_%", escape="\\"),
        )
        .distinct()
    )
    return {t for (t,) in rows}


def _price_on(db, ticker: str, d: datetime.date) -> float | None:
    row = (
        db.query(PricesDaily.close)
        .filter(PricesDaily.ticker == ticker, PricesDaily.trade_date <= d)
        .order_by(PricesDaily.trade_date.desc())
        .first()
    )
    return float(row[0]) if row and row[0] else None


def _latest_quarter(db, ticker: str) -> str:
    r = (
        db.query(FinancialsQuarterly.fiscal_year, FinancialsQuarterly.fiscal_quarter)
        .filter(
            FinancialsQuarterly.ticker == ticker, FinancialsQuarterly.fiscal_quarter > 0
        )
        .order_by(
            FinancialsQuarterly.fiscal_year.desc(),
            FinancialsQuarterly.fiscal_quarter.desc(),
        )
        .first()
    )
    return f"Q{r[1]}/{r[0]}" if r else ""


def _reports(db, ticker: str, as_of: datetime.date) -> list:
    return (
        db.query(Consensus)
        .filter(
            Consensus.ticker == ticker,
            Consensus.report_date >= as_of - datetime.timedelta(days=WINDOW_DAYS),
            Consensus.report_date <= as_of,
            Consensus.target_price > 0,
            Consensus.is_synthetic.isnot(True),
        )
        .order_by(Consensus.report_date.desc())
        .all()
    )


def _pct(x: float) -> str:
    return f"{x:+.0%}".replace("%", "%")


def collect(db, as_of: datetime.date) -> dict:
    with open(PROJECT_ROOT / "config" / "analyst_notes.yaml", encoding="utf-8") as fh:
        notes_cfg = yaml.safe_load(fh) or {}
    repair = _repair_status(_fixed_tickers(db))
    rows: list[dict] = []
    reports: list[dict] = []
    assumptions: list[dict] = []
    fin: list[dict] = []
    flag_seen: dict[str, dict] = {}

    for t in get_vn100_symbols():
        out = value_ticker(db, t)
        meta = db.query(Ticker).filter(Ticker.ticker == t).one_or_none()
        plan = route(t) or {}
        price_date = (
            db.query(func.max(PricesDaily.trade_date))
            .filter(PricesDaily.ticker == t)
            .scalar()
        )
        # Company | CompanyBank: nhánh theo is_bank lúc chạy nên để Any.
        company: Any = None
        res: dict = {}
        err = out.get("error")
        try:
            company = build_company_data(db, t, mode="TTM", fetch_live=False)
            res = valuate(company)
        # Lỗi 1 mã hiện ở cột trạng thái, không chặn cả bảng.
        except Exception as e:  # noqa: BLE001
            err = err or f"{type(e).__name__}: {str(e)[:80]}"

        flags = list(out.get("flags") or [])
        for d in describe_flags(flags):
            flag_seen.setdefault(str(d["code"]).split(":")[0], d)
        rated = out.get("fair_value") is not None
        view = get_consensus_view(db, t, as_of=as_of, window_days=WINDOW_DAYS)
        used = {(q.broker_raw, q.report_date, q.target_price) for q in view.quotes}
        price = out.get("price") or 0.0

        # --- ghi chú tự động (chỉ từ dữ liệu) ---
        auto: list[str] = []
        move = None
        if err:
            status = f"Lỗi: {err}"
        elif rated:
            status = "Công bố"
        else:
            reasons = [_BLOCK_REASON[f] for f in flags if f in _BLOCK_REASON]
            reasons += [
                c["message"]
                for c in res.get("data_checks", [])
                if c.get("severity") == "error"
            ]
            status = "Không xếp hạng: " + (
                "; ".join(dict.fromkeys(reasons)) or "cổng phát hành"
            )
            if res.get("blended_fair_value_per_share"):
                auto.append(
                    f"FV chẩn đoán nội bộ {res['blended_fair_value_per_share']:,.0f}đ — KHÔNG dùng để ra quyết định."
                )

        if not view.has_data:
            auto.append(f"Không có báo cáo CTCK trong {WINDOW_DAYS} ngày.")
        else:
            if view.thin:
                auto.append("Chỉ 1 CTCK theo dõi — không phải đồng thuận thị trường.")
            if view.stale:
                auto.append(
                    f"Báo cáo CTCK mới nhất đã {view.newest_age_days} ngày (> {DEFAULT_STALE_AFTER_DAYS})."
                )
            pxs = [p for q in view.quotes if (p := _price_on(db, t, q.report_date))]
            if pxs and price:
                move = price / statistics.median(pxs) - 1
                if abs(move) > PRICE_MOVE_NOTE:
                    auto.append(
                        f"Giá đã {'tăng' if move > 0 else 'giảm'} {abs(move):.0%} so với lúc CTCK ra báo cáo — giá mục tiêu có thể lỗi thời."
                    )
            if rated and view.median:
                dev = out["fair_value"] / view.median - 1
                if abs(dev) > BAND:
                    auto.append(
                        f"GTHL lệch {_pct(dev)} so với trung vị CTCK (ngưỡng ±{BAND:.0%}) — cần giải thích, không ép khớp."
                    )

        ia, ra = res.get("intrinsic_fv"), res.get("relative_fv")
        if ia and ra and min(ia, ra) > 0 and max(ia, ra) / min(ia, ra) > 1.5:
            auto.append(
                f"Hai phương pháp lệch lớn: nội tại {ia:,.0f} vs tương đối {ra:,.0f}."
            )

        is_bank = company is not None and hasattr(
            company.historical_is[-1], "net_interest_income"
        )
        a: Any = company.assumptions if company is not None else None
        bank_extra = {}
        if is_bank:
            is_, bs = company.historical_is, company.historical_bs
            loans_avg = (
                (bs[-2].customer_loans + bs[-1].customer_loans) / 2
                if len(bs) > 1
                else 0
            )
            cc_ttm = is_[-1].provision_expense / loans_avg if loans_avg else None
            ea_now = (
                bs[-1].total_equity / bs[-1].total_assets
                if bs[-1].total_assets
                else None
            )
            proj = res.get("projections") or []
            ni1 = proj[0]["net_income"] if proj else None
            ni_ttm = is_[-1].net_income
            bank_extra = {
                "cc_ttm": cc_ttm,
                "ea_now": ea_now,
                "ni1": ni1,
                "ni_ttm": ni_ttm,
                "roe_ttm": (
                    ni_ttm / bs[-1].total_equity if bs[-1].total_equity else None
                ),
            }
            if ni1 is not None and ni_ttm > 0 and ni1 < BANK_NI_DROP_NOTE * ni_ttm:
                why = []
                cc_n = a.credit_cost[0]
                if cc_ttm and cc_n > cc_ttm * 1.2:
                    why.append(
                        f"chi phí tín dụng chuẩn hóa {cc_n:.2%} (trung vị 3 năm tài chính đầy đủ) so với TTM ~{cc_ttm:.2%}"
                    )
                if (
                    a.min_equity_to_assets
                    and ea_now
                    and ea_now < a.min_equity_to_assets
                ):
                    why.append(
                        f"ràng buộc vốn E/A tối thiểu {a.min_equity_to_assets:.1%} (trung vị lịch sử) > E/A hiện tại {ea_now:.1%} nên tài sản bị ép co lại năm 1"
                    )
                auto.append(
                    f"Dự phóng LNST năm 1 {ni1:,.0f} tỷ thấp hơn TTM {ni_ttm:,.0f} tỷ ({ni1 / ni_ttm - 1:+.0%})"
                    + (" do " + "; ".join(why) if why else "")
                    + ". Giả định thận trọng của mô hình — chờ duyệt."
                )

        dq = []
        if any(
            f.startswith("UNEXPLAINED_CURRENT_ASSETS")
            for f in flags + list(getattr(company, "data_flags", []) or [])
        ):
            dq.append(
                "Tài sản ngắn hạn có phần dư lớn chưa giải thích (nghi mất dòng khi nạp)"
            )
        if repair.get(t, "").startswith("Chưa"):
            dq.append(
                "BCTC chưa sửa được dòng trùng item_id (lỗi mạng) — số dư tài sản có thể thiếu"
            )
        if dq:
            auto.append("Dữ liệu: " + "; ".join(dq) + ".")

        note = " ".join(filter(None, [notes_cfg.get(t, ""), *auto]))
        caution = " ".join(
            d["message"].split(". ")[0].rstrip(".") + "."
            for d in describe_flags(
                [f for f in flags if f not in ("NOT_RATEABLE", "DATA_REVIEW_REQUIRED")]
            )
            if d["level"] in ("error", "warning")
        )
        titles = [
            f"{r.report_date:%d/%m} {r.broker_canon or r.broker}: {r.report_title}"
            for r in _reports(db, t, as_of)
            if r.report_title
        ][:3]

        w = res.get("weight_intrinsic")
        method = _METHOD_LABEL.get(out.get("method", ""), out.get("method", ""))
        if w is not None:
            method += f" (nội tại {w:.0%} / tương đối {1 - w:.0%})"
        rows.append(
            {
                "ticker": t,
                "name": meta.company_name if meta else "",
                "exchange": meta.exchange if meta else "",
                "sector": plan.get("group") or (meta.sector if meta else ""),
                "nature": (res.get("decision") or {}).get("business_nature", ""),
                "method": method,
                "broker_method": " + ".join(filter(None, [plan.get("raw_main")])),
                "price_date": price_date,
                "price": price or None,
                "intrinsic": ia,
                "relative": ra,
                "fv": out.get("fair_value"),
                "mos": (res.get("decision") or {}).get("target_mos"),
                "engine_rec": REC_LABELS_VI.get(
                    out.get("recommendation", ""), out.get("recommendation", "")
                ),
                "status": status,
                "cons_median": view.median,
                "px_move": move,
                "caution": caution,
                "note": note,
                "titles": "\n".join(titles),
                "flags": ", ".join(flags),
                "rated": rated and not err,
                "upside_sort": out.get("upside"),
            }
        )

        for r in _reports(db, t, as_of):
            reports.append(
                {
                    "ticker": t,
                    "broker": r.broker_canon or r.broker,
                    "date": r.report_date,
                    "tp": float(r.target_price),
                    "rating": r.rating or "",
                    "px": _price_on(db, t, r.report_date),
                    "used": (
                        "Có"
                        if (str(r.broker), r.report_date, float(r.target_price)) in used
                        else "Không"
                    ),
                    "age": (as_of - r.report_date).days,
                    "title": r.report_title or "",
                    "site": r.source_site or "",
                    "url": r.source_url or "",
                }
            )

        if company is None:
            continue
        is_, bs = company.historical_is, company.historical_bs
        fy = is_[:-1]  # phần tử cuối là TTM
        rev = (lambda x: x.total_operating_income) if is_bank else (lambda x: x.revenue)
        eq = bs[-1].total_equity if is_bank else bs[-1].parent_equity
        fin.append(
            {
                "ticker": t,
                "kind": "Ngân hàng" if is_bank else "Phi ngân hàng",
                "shares": company.shares_outstanding,
                "period": _latest_quarter(db, t),
                "rev_ttm": rev(is_[-1]),
                "ni_ttm": is_[-1].net_income,
                "equity": eq,
                "debt": (
                    None if is_bank else bs[-1].short_term_debt + bs[-1].long_term_debt
                ),
                "fy0_year": fy[-4].year if len(fy) >= 4 else None,
                "fy0_rev": rev(fy[-4]) if len(fy) >= 4 else None,
                "fy_year": fy[-1].year if fy else None,
                "fy_rev": rev(fy[-1]) if fy else None,
                "fy_ni": fy[-1].net_income if fy else None,
                "fy_prev_ni": fy[-2].net_income if len(fy) >= 2 else None,
                "repair": repair.get(t, "Không cần/không phát hiện"),
                "dq_flags": ", ".join(
                    f
                    for f in flags + list(company.data_flags or [])
                    if f.split(":")[0]
                    in (
                        "UNEXPLAINED_CURRENT_ASSETS",
                        "STALE_PRICE",
                        "STALE_MACRO_RF",
                        "SOURCE_CONFLICT",
                        "MISSING_FINANCIAL_INPUT",
                        "DATA_INCOMPLETE",
                        "FINANCIAL_QC_MISSING",
                    )
                ),
            }
        )
        g = getattr(a, "terminal_growth_rate", None)
        base = {
            "ticker": t,
            "kind": fin[-1]["kind"],
            "rf": a.risk_free_rate,
            "beta": a.beta,
            "erp": a.erp,
            "coe_override": a.cost_of_equity,
            "g": g,
            "pb_hist": a.historical_pb_median,
        }
        if is_bank:
            base.update(
                cg1=a.credit_growth[0],
                nim=a.nim[0],
                cir=a.cir[0],
                cc=a.credit_cost[0],
                cc_ttm=bank_extra["cc_ttm"],
                roe_ttm=bank_extra["roe_ttm"],
                # Cột 'ROE bền vững (terminal)': ROE năm 5 sau hội tụ ROE ngành (λ), trước trần.
                roe_s=converged_roe(a.sustainable_roe, a.industry_roe, a.roe_convergence),
                payout=a.dividend_payout_ratio,
                ea_min=a.min_equity_to_assets,
                ea_now=bank_extra["ea_now"],
                ni1=bank_extra["ni1"],
                ni_ttm=bank_extra["ni_ttm"],
            )
        else:
            base.update(
                kd=a.cost_of_debt,
                tax=a.tax_rate,
                g1=a.revenue_growth[0],
                ebit_m=a.ebit_margin[0],
                capex=a.capex_to_revenue[0],
                pe=a.target_pe,
                ev=a.target_ev_ebitda,
            )
        assumptions.append(base)

    rows.sort(
        key=lambda r: (
            not r["rated"],
            -(r["upside_sort"] or 0) if r["rated"] else 0,
            r["ticker"],
        )
    )
    reports.sort(key=lambda r: (r["ticker"], r["date"]), reverse=False)
    reports.sort(key=lambda r: r["date"], reverse=True)
    reports.sort(key=lambda r: r["ticker"])
    return {
        "rows": rows,
        "reports": reports,
        "assumptions": assumptions,
        "fin": fin,
        "flags": flag_seen,
    }


# ---------------------------------------------------------------- ghi Excel
def _header(ws, row: int, headers: list[tuple[str, int]]) -> None:
    for i, (h, width) in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.alignment = (
            F_HEAD,
            FILL_HEAD,
            Alignment(wrap_text=True, vertical="center"),
        )
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 42


def _put(
    ws,
    r: int,
    c: int,
    v,
    fmt: str | None = None,
    font: Font = F_BASE,
    wrap: bool = False,
):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if wrap:
        cell.alignment = WRAP
    return cell


def write_book(data: dict, as_of: datetime.date, path: Path) -> None:
    wb = Workbook()
    readme = wb.active
    readme.title = "Hướng dẫn"
    ws = wb.create_sheet("Tổng hợp")
    wr = wb.create_sheet("Giá mục tiêu CTCK")
    wa = wb.create_sheet("Giả định")
    wf = wb.create_sheet("Số liệu TC")
    wm = wb.create_sheet("Phương pháp")
    wg = wb.create_sheet("Chú thích cờ")
    wq = wb.create_sheet("Chất lượng dữ liệu")

    # ---------- Phương pháp (tham số toàn cục được công thức tham chiếu) ----------
    wm["A1"], wm["A1"].font = "Phương pháp định giá & tham số chung", F_TITLE
    params = [
        (
            "Lãi suất phi rủi ro rf (TPCP VN 10 năm, động)",
            data["assumptions"][0]["rf"] if data["assumptions"] else None,
            FMT_PCT,
            "Nguồn: macro_series TPCP_10Y trong DB; dùng chung cho mọi mã.",
        ),
        (
            "ERP = ERP thị trường phát triển 4,5% + CRP Việt Nam 3,7%",
            float(CFG["coe_convention"]["erp_total"]),
            FMT_PCT,
            "Damodaran (convention VND-base) — config/defaults.yaml: coe_convention.",
        ),
        (
            "Tăng trưởng vĩnh viễn g = min(4%, rf, GDP danh nghĩa 8%)",
            float(CFG["coe_convention"]["terminal_growth"]),
            FMT_PCT,
            "Damodaran: g ≤ rf. config/defaults.yaml: coe_convention.terminal_growth.",
        ),
        (
            "Ngưỡng GIẢM TỶ TRỌNG (upside dưới 0 nhưng ≥ ngưỡng này)",
            TRIM_FLOOR,
            FMT_PCT,
            "config/defaults.yaml: margin_of_safety.trim_floor. Dưới ngưỡng → BÁN.",
        ),
        (
            "Ngưỡng lệch cần giải thích so với trung vị CTCK",
            BAND,
            FMT_PCT,
            "config/defaults.yaml: audit.consensus_review_band.",
        ),
    ]
    for i, (label, val, fmt, src) in enumerate(params, 3):
        _put(wm, i, 1, label, font=F_BOLD)
        c = _put(wm, i, 2, val, fmt, F_INPUT)
        c.fill = FILL_KEY
        _put(wm, i, 3, src, wrap=True)
    trim_ref = "'Phương pháp'!$B$6"
    r = len(params) + 4
    _put(wm, r, 1, "Biên an toàn (MOS) yêu cầu để MUA theo bản chất ngành", font=F_BOLD)
    for nature, mos in (CFG["margin_of_safety"]["by_business_nature"] or {}).items():
        r += 1
        _put(wm, r, 1, nature)
        _put(wm, r, 2, float(mos), FMT_PCT, F_INPUT)
    r += 1
    _put(wm, r, 1, "Mặc định / cộng thêm khi vĩ mô căng thẳng")
    _put(
        wm,
        r,
        2,
        f"{CFG['margin_of_safety']['default']:.0%} / +{CFG['margin_of_safety']['macro_stress_add']:.0%}",
    )
    r += 2
    methods = [
        (
            "Thu nhập thặng dư + P/B (ngân hàng)",
            (
                "RI = BVPS + PV(ROE − COE) × vốn CSH đầu kỳ, dự phóng 5 năm từ NIM, tăng trưởng tín dụng, CIR, chi phí tín dụng; "
                "giá trị cuối kỳ Gordon với ROE bền vững (trần 20%). Chân tương đối: P/B trung vị lịch sử của chính NH × BVPS. "
                "Không dùng EV/EBITDA hay DCF cho ngân hàng (nợ là nguyên liệu sản xuất)."
            ),
        ),
        (
            "P/B (chứng khoán, bảo hiểm)",
            "P/B trung vị lịch sử của chính mã (đã lọc ngoại lai) × BVPS hiện tại, có sàn/trần ngành. CTCK bám chu kỳ thanh khoản.",
        ),
        (
            "DCF/FCFF (phi tài chính)",
            (
                "FCFF 5 năm = NOPAT + khấu hao − capex − Δvốn lưu động; chiết khấu WACC (COE CAPM, Kd sau thuế); terminal Gordon g. "
                "Cầu nối EV → vốn CSH: − nợ vay + tiền & ĐTNH + đầu tư dài hạn − CĐKKS. Biên EBIT chuẩn hóa giữa chu kỳ cho ngành chu kỳ."
            ),
        ),
        (
            "EV/EBITDA (chuẩn hóa)",
            "EBITDA giữa chu kỳ × EV/EBITDA trung vị lịch sử của chính mã (lọc ngoại lai) → trừ nợ ròng.",
        ),
        (
            "P/E (EPS chuẩn hóa)",
            "EPS cổ đông mẹ chuẩn hóa × P/E trung vị lịch sử của chính mã (bỏ năm LN đáy/đột biến). Không dùng cho DN lỗ.",
        ),
        (
            "DDM (điện)",
            "Cổ tức dự phóng chiết khấu COE; blend với DCF (DCF chủ đạo) cho nhóm điện.",
        ),
        (
            "RNAV / SOTP (BĐS, tập đoàn)",
            "Cần danh mục dự án/mảng kinh doanh được người duyệt. Khi chưa có → chỉ là proxy → KHÔNG xếp hạng (không bịa số).",
        ),
        (
            "Khuyến nghị",
            (
                "MUA khi upside ≥ MOS theo bản chất ngành; NẮM GIỮ khi 0 ≤ upside < MOS; GIẢM TỶ TRỌNG khi ngưỡng ≤ upside < 0; còn lại BÁN. "
                "Cổng phát hành chặn mọi khuyến nghị khi mô hình là proxy, kết quả cực đoan (>+150% / <−60%), dữ liệu lỗi."
            ),
        ),
        (
            "Giá mục tiêu CTCK",
            (
                "Nguồn Simplize/24hMoney, cửa sổ 180 ngày, mỗi CTCK chỉ lấy báo cáo mới nhất (cột 'Dùng trong trung vị'). "
                "Giá mục tiêu và giá lịch sử VCI đều đã điều chỉnh quyền (kiểm tra 44 sự kiện chia/thưởng 2026: không có khoảng trống giá) — không điều chỉnh thêm."
            ),
        ),
    ]
    _put(wm, r, 1, "Phương pháp", font=F_BOLD)
    _put(wm, r, 2, "Cách tính", font=F_BOLD)
    for name, how in methods:
        r += 1
        _put(wm, r, 1, name, font=F_BOLD, wrap=True)
        _put(wm, r, 2, how, wrap=True)
        wm.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        wm.row_dimensions[r].height = 58
    wm.column_dimensions["A"].width = 52
    wm.column_dimensions["B"].width = 60
    wm.column_dimensions["C"].width = 60

    # ---------- Số liệu TC ----------
    fin_cols = [
        ("Mã", 8),
        ("Loại", 14),
        ("Số CP lưu hành (triệu)", 14),
        ("Kỳ BCTC mới nhất", 11),
        ("DT/TOI TTM (tỷ)", 14),
        ("LNST cổ đông mẹ TTM (tỷ)", 14),
        ("VCSH mẹ (tỷ)", 14),
        ("Nợ vay (tỷ)", 14),
        ("Năm gốc CAGR", 9),
        ("DT/TOI năm gốc (tỷ)", 14),
        ("Năm FY gần nhất", 9),
        ("DT/TOI FY gần nhất (tỷ)", 14),
        ("LNST FY gần nhất (tỷ)", 14),
        ("LNST FY trước (tỷ)", 14),
    ]
    _header(wf, 1, fin_cols)
    for i, f in enumerate(data["fin"], 2):
        vals = [
            f["ticker"],
            f["kind"],
            f["shares"],
            f["period"],
            f["rev_ttm"],
            f["ni_ttm"],
            f["equity"],
            f["debt"],
            str(f["fy0_year"]) if f["fy0_year"] else None,
            f["fy0_rev"],
            str(f["fy_year"]) if f["fy_year"] else None,
            f["fy_rev"],
            f["fy_ni"],
            f["fy_prev_ni"],
        ]
        for j, v in enumerate(vals, 1):
            _put(
                wf,
                i,
                j,
                v,
                FMT_VND if isinstance(v, float) else None,
                F_INPUT if j > 2 else F_BASE,
            )
    nfin = len(data["fin"]) + 1
    wf.freeze_panes = "B2"
    wf.auto_filter.ref = f"A1:{get_column_letter(len(fin_cols))}{nfin}"
    _put(
        wf,
        nfin + 2,
        1,
        "Nguồn: financials_quarterly (vnstock/VCI, BCTC hợp nhất), đơn vị tỷ đồng. TTM = 4 quý gần nhất; "
        "DT/TOI = doanh thu thuần (phi NH) hoặc tổng thu nhập hoạt động (NH). VCSH mẹ = VCSH − CĐKKS.",
        wrap=False,
    )

    def fin_ref(col: int, key: str) -> str:
        L = get_column_letter(col)
        return f"INDEX('Số liệu TC'!${L}$2:${L}${nfin},MATCH({key},'Số liệu TC'!$A$2:$A${nfin},0))"

    # ---------- Giá mục tiêu CTCK ----------
    rep_cols = [
        ("Mã", 8),
        ("CTCK", 12),
        ("Ngày báo cáo", 12),
        ("Giá mục tiêu (VND)", 13),
        ("Khuyến nghị CTCK", 16),
        ("Giá TT ngày BC (VND)", 13),
        ("Upside tại ngày BC", 11),
        ("Giá TT hiện tại (VND)", 13),
        ("Upside còn lại", 11),
        ("Dùng trong trung vị", 10),
        ("Tuổi BC (ngày)", 9),
        ("Tiêu đề báo cáo", 70),
        ("Nguồn", 11),
        ("URL", 40),
    ]
    _header(wr, 1, rep_cols)
    nsum = len(data["rows"]) + 1
    for i, p in enumerate(data["reports"], 2):
        _put(wr, i, 1, p["ticker"])
        _put(wr, i, 2, p["broker"])
        _put(wr, i, 3, p["date"], FMT_DATE)
        _put(wr, i, 4, p["tp"], FMT_VND, F_INPUT)
        _put(wr, i, 5, p["rating"])
        _put(wr, i, 6, p["px"], FMT_VND, F_INPUT)
        _put(wr, i, 7, f'=IF(AND(ISNUMBER(F{i}),F{i}>0),D{i}/F{i}-1,"")', FMT_PCT)
        _put(
            wr,
            i,
            8,
            f"=INDEX('Tổng hợp'!$H$2:$H${nsum},MATCH(A{i},'Tổng hợp'!$A$2:$A${nsum},0))",
            FMT_VND,
            F_LINK,
        )
        _put(wr, i, 9, f'=IF(AND(ISNUMBER(H{i}),H{i}>0),D{i}/H{i}-1,"")', FMT_PCT)
        _put(wr, i, 10, p["used"])
        _put(wr, i, 11, p["age"])
        _put(wr, i, 12, p["title"])
        _put(wr, i, 13, p["site"])
        _put(wr, i, 14, p["url"])
    nrep = len(data["reports"]) + 1
    wr.freeze_panes = "C2"
    wr.auto_filter.ref = f"A1:N{nrep}"

    # ---------- Tổng hợp ----------
    # (khóa, tiêu đề, độ rộng) — công thức tham chiếu cột theo khóa, thêm/bớt cột không lệch chữ cái.
    spec = [
        ("tk", "Mã", 7),
        ("name", "Tên công ty", 30),
        ("ex", "Sàn", 7),
        ("sec", "Ngành", 14),
        ("nat", "Bản chất KD", 12),
        ("meth", "Phương pháp định giá", 30),
        ("pd", "Ngày giá", 11),
        ("px", "Giá thị trường (VND)", 12),
        ("int", "GTHL nội tại (VND)", 12),
        ("rel", "GTHL tương đối (VND)", 12),
        ("fv", "GIÁ TRỊ HỢP LÝ (VND)", 13),
        ("up", "Upside", 9),
        ("mos", "MOS yêu cầu", 8),
        ("rec", "Khuyến nghị", 13),
        ("conf", "Độ tin cậy", 22),
        ("st", "Trạng thái công bố", 34),
        ("med", "Trung vị giá MT CTCK (VND)", 13),
        ("n", "Số CTCK", 7),
        ("lo", "Giá MT thấp nhất", 11),
        ("hi", "Giá MT cao nhất", 11),
        ("last", "BC CTCK mới nhất", 11),
        ("mv", "Giá thay đổi từ lúc CTCK ra BC", 10),
        ("upc", "Upside theo CTCK", 9),
        ("dev", "Lệch GTHL vs CTCK", 9),
        ("mcap", "Vốn hóa (tỷ)", 12),
        ("pe", "P/E TTM", 8),
        ("pb", "P/B", 7),
        ("roe", "ROE TTM", 8),
        ("nm", "Biên LN ròng", 8),
        ("de", "Nợ vay/VCSH", 8),
        ("cagr", "Tăng trưởng DT 3 năm (CAGR)", 10),
        ("nig", "Tăng trưởng LNST FY gần nhất", 10),
        ("caut", "Chú ý (cờ mô hình)", 50),
        ("note", "Lưu ý định giá", 70),
        ("tit", "Luận điểm CTCK gần nhất (tiêu đề BC)", 60),
        ("flags", "Mã cờ hệ thống", 30),
    ]
    C = {k: get_column_letter(j) for j, (k, _, _) in enumerate(spec, 1)}
    N = {k: j for j, (k, _, _) in enumerate(spec, 1)}
    _header(ws, 1, [(h, w) for _, h, w in spec])
    band_ref = "'Phương pháp'!$B$7"
    rep_rng = lambda L: f"'Giá mục tiêu CTCK'!${L}$2:${L}${nrep}"
    fin_tickers = {f["ticker"] for f in data["fin"]}
    for i, x in enumerate(data["rows"], 2):
        c_ = {k: f"{L}{i}" for k, L in C.items()}
        key = f"${C['tk']}{i}"
        put = lambda k, v, fmt=None, font=F_BASE, wrap=False, r=i: _put(
            ws, r, N[k], v, fmt, font, wrap
        )
        put("tk", x["ticker"], font=F_BOLD)
        put("name", x["name"])
        put("ex", x["exchange"])
        put("sec", x["sector"])
        put("nat", x["nature"])
        put("meth", x["method"], wrap=True)
        put("pd", x["price_date"], FMT_DATE)
        put("px", x["price"], FMT_VND, F_INPUT)
        put("int", x["intrinsic"] if x["rated"] else None, FMT_VND, F_INPUT)
        put("rel", x["relative"] if x["rated"] else None, FMT_VND, F_INPUT)
        put(
            "fv", x["fv"], FMT_VND, Font(name=ARIAL, size=10, bold=True, color="0000FF")
        ).fill = FILL_KEY
        put(
            "up",
            f'=IF(AND(ISNUMBER({c_["fv"]}),{c_["px"]}>0),{c_["fv"]}/{c_["px"]}-1,"")',
            FMT_PCT,
        )
        put("mos", x["mos"], FMT_PCT, F_INPUT)
        put(
            "rec",
            f'=IF({c_["st"]}<>"Công bố","{REC_LABELS_VI["NOT_RATED"]}",IF({c_["up"]}>={c_["mos"]},"MUA",'
            f'IF({c_["up"]}>=0,"NẮM GIỮ",IF({c_["up"]}>={trim_ref},"GIẢM TỶ TRỌNG","BÁN"))))',
            font=F_BOLD,
        )
        put(
            "conf",
            f'=IF({c_["st"]}<>"Công bố","—",IF({c_["n"]}<2,"Chưa đối chiếu được (≤1 CTCK)",'
            f'IF(AND(ISNUMBER({c_["mv"]}),ABS(N({c_["mv"]}))>{band_ref}),'
            f'"Chưa đối chiếu được (giá biến động mạnh sau BC CTCK)",'
            f'IF(ABS({c_["dev"]})>{band_ref},"THẤP: lệch CTCK > ±"&TEXT({band_ref},"0%"),'
            f'"Trung bình: trong ±"&TEXT({band_ref},"0%")&" CTCK"))))',
        )
        put("st", x["status"], wrap=True)
        put("med", x["cons_median"], FMT_VND, F_INPUT)
        put("n", f'=COUNTIFS({rep_rng("A")},{key},{rep_rng("J")},"Có")', "0", F_LINK)
        put(
            "lo",
            f'=IF({c_["n"]}>0,_xlfn.MINIFS({rep_rng("D")},{rep_rng("A")},{key},{rep_rng("J")},"Có"),"")',
            FMT_VND,
            F_LINK,
        )
        put(
            "hi",
            f'=IF({c_["n"]}>0,_xlfn.MAXIFS({rep_rng("D")},{rep_rng("A")},{key},{rep_rng("J")},"Có"),"")',
            FMT_VND,
            F_LINK,
        )
        put(
            "last",
            f'=IF({c_["n"]}>0,_xlfn.MAXIFS({rep_rng("C")},{rep_rng("A")},{key}),"")',
            FMT_DATE,
            F_LINK,
        )
        put("mv", x["px_move"], FMT_PCT, F_INPUT)
        put(
            "upc",
            f'=IF(AND(ISNUMBER({c_["med"]}),{c_["px"]}>0),{c_["med"]}/{c_["px"]}-1,"")',
            FMT_PCT,
        )
        put(
            "dev",
            f'=IF(AND(ISNUMBER({c_["fv"]}),ISNUMBER({c_["med"]})),{c_["fv"]}/{c_["med"]}-1,"")',
            FMT_PCT,
        )
        if x["ticker"] in fin_tickers:
            sh, rev, ni, eq, debt = (
                fin_ref(3, key),
                fin_ref(5, key),
                fin_ref(6, key),
                fin_ref(7, key),
                fin_ref(8, key),
            )
            fy0, fy0y, fyr, fyy = (
                fin_ref(10, key),
                fin_ref(9, key),
                fin_ref(12, key),
                fin_ref(11, key),
            )
            fyni, fyp = fin_ref(13, key), fin_ref(14, key)
            put("mcap", f"={c_['px']}*{sh}/1000", FMT_VND, F_LINK)
            put("pe", f'=IFERROR(IF({ni}>0,{c_["mcap"]}/{ni},""),"")', FMT_X, F_LINK)
            put("pb", f'=IFERROR(IF({eq}>0,{c_["mcap"]}/{eq},""),"")', FMT_X, F_LINK)
            put("roe", f'=IFERROR(IF({eq}>0,{ni}/{eq},""),"")', FMT_PCT, F_LINK)
            put("nm", f'=IFERROR(IF({rev}>0,{ni}/{rev},""),"")', FMT_PCT, F_LINK)
            put(
                "de",
                f'=IFERROR(IF(AND(ISNUMBER({debt}),{eq}>0),{debt}/{eq},""),"")',
                FMT_X,
                F_LINK,
            )
            put(
                "cagr",
                f'=IFERROR(IF(AND({fy0}>0,{fyr}>0),({fyr}/{fy0})^(1/({fyy}-{fy0y}))-1,""),"")',
                FMT_PCT,
                F_LINK,
            )
            put("nig", f'=IFERROR(IF({fyp}>0,{fyni}/{fyp}-1,""),"")', FMT_PCT, F_LINK)
        put("caut", x["caution"], wrap=True)
        put("note", x["note"], wrap=True)
        put("tit", x["titles"], wrap=True)
        put("flags", x["flags"], wrap=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spec))}{nsum}"
    ws[f"{C['fv']}1"].comment = Comment(
        "Đầu ra engine định giá (batch.value_ticker = cùng lõi Streamlit). Để trống = không xếp hạng. "
        "Có thể sửa: Upside/Khuyến nghị/Lệch CTCK tự tính lại.",
        "VN100",
    )
    ws[f"{C['med']}1"].comment = Comment(
        "consensus_view: báo cáo 180 ngày, mỗi CTCK 1 báo cáo mới nhất, bỏ dữ liệu synthetic.",
        "VN100",
    )
    ws[f"{C['mos']}1"].comment = Comment(
        "Biên an toàn yêu cầu để MUA theo bản chất ngành (decision_engine), xem sheet Phương pháp.",
        "VN100",
    )
    ws[f"{C['conf']}1"].comment = Comment(
        "Chưa mô hình nào có golden test đầu-cuối (luật vàng #8). 'THẤP' = lệch trung vị CTCK quá ngưỡng: "
        "khuyến nghị chỉ mang tính tham khảo cho tới khi có giải thích/duyệt.",
        "VN100",
    )

    # ---------- Giả định ----------
    as_cols = [
        ("Mã", 7),
        ("Loại", 13),
        ("rf", 7),
        ("Beta", 7),
        ("ERP", 7),
        ("COE", 7),
        ("g vĩnh viễn", 8),
        ("P/B lịch sử (trung vị)", 9),
        ("Kd trước thuế", 8),
        ("Thuế TNDN", 8),
        ("Tăng trưởng DT năm 1", 9),
        ("Biên EBIT", 8),
        ("Capex/DT", 8),
        ("P/E mục tiêu", 8),
        ("EV/EBITDA mục tiêu", 9),
        ("Tăng trưởng tín dụng năm 1", 9),
        ("NIM", 7),
        ("CIR", 7),
        ("CP tín dụng chuẩn hóa", 9),
        ("CP tín dụng TTM (xấp xỉ)", 9),
        ("ROE TTM", 8),
        ("ROE bền vững (terminal)", 9),
        ("Payout tiền mặt", 8),
        ("E/A tối thiểu (mô hình)", 9),
        ("E/A hiện tại", 8),
        ("LNST dự phóng năm 1 (tỷ)", 12),
        ("LNST TTM (tỷ)", 12),
    ]
    _header(wa, 1, as_cols)
    keys = [
        None,
        None,
        "rf",
        "beta",
        "erp",
        None,
        "g",
        "pb_hist",
        "kd",
        "tax",
        "g1",
        "ebit_m",
        "capex",
        "pe",
        "ev",
        "cg1",
        "nim",
        "cir",
        "cc",
        "cc_ttm",
        "roe_ttm",
        "roe_s",
        "payout",
        "ea_min",
        "ea_now",
        "ni1",
        "ni_ttm",
    ]
    fmts = (
        [None, None]
        + [
            FMT_PCT,
            "0.00",
            FMT_PCT,
            FMT_PCT,
            FMT_PCT,
            FMT_X,
            FMT_PCT,
            FMT_PCT,
            FMT_PCT,
            FMT_PCT,
            FMT_PCT,
            FMT_X,
            FMT_X,
        ]
        + [FMT_PCT] * 10
        + [FMT_VND, FMT_VND]
    )
    for i, a in enumerate(data["assumptions"], 2):
        _put(wa, i, 1, a["ticker"], font=F_BOLD)
        _put(wa, i, 2, a["kind"])
        for j, k in enumerate(keys, 1):
            if k is None:
                continue
            _put(wa, i, j, a.get(k), fmts[j - 1], F_INPUT)
        if a.get("coe_override"):
            _put(wa, i, 6, a["coe_override"], FMT_PCT, F_INPUT)
        else:
            _put(wa, i, 6, f"=C{i}+D{i}*E{i}", FMT_PCT)
    na = len(data["assumptions"]) + 1
    wa.freeze_panes = "C2"
    wa.auto_filter.ref = f"A1:{get_column_letter(len(as_cols))}{na}"
    _put(
        wa,
        na + 2,
        1,
        "Giả định base case do máy đặt (build_company_data) — chỉ là điểm khởi đầu, chỉnh trong app Streamlit. "
        "COE = rf + Beta × ERP (Damodaran, VND-base). CP tín dụng chuẩn hóa = trung vị 3 năm tài chính đầy đủ; "
        "E/A tối thiểu = trung vị 3 năm gần nhất (ràng buộc vốn thay cho CAR).",
    )

    # ---------- Chú thích cờ ----------
    _header(wg, 1, [("Mã cờ", 34), ("Mức độ", 10), ("Giải thích", 110)])
    for i, (code, d) in enumerate(sorted(data["flags"].items()), 2):
        _put(wg, i, 1, code)
        _put(
            wg,
            i,
            2,
            {"error": "Nghiêm trọng", "warning": "Cảnh báo", "info": "Thông tin"}.get(
                d["level"], d["level"]
            ),
        )
        _put(wg, i, 3, d["message"], wrap=True)

    # ---------- Chất lượng dữ liệu ----------
    _header(
        wq,
        1,
        [
            ("Mã", 8),
            ("Ngày giá mới nhất", 12),
            ("Kỳ BCTC mới nhất", 12),
            ("Sửa dòng trùng item_id (BCTC)", 30),
            ("Cờ dữ liệu", 50),
        ],
    )
    pdate = {x["ticker"]: x["price_date"] for x in data["rows"]}
    for i, f in enumerate(sorted(data["fin"], key=lambda f: f["ticker"]), 2):
        _put(wq, i, 1, f["ticker"], font=F_BOLD)
        _put(wq, i, 2, pdate.get(f["ticker"]), FMT_DATE)
        _put(wq, i, 3, f["period"])
        _put(wq, i, 4, f["repair"])
        _put(wq, i, 5, f["dq_flags"], wrap=True)
    wq.freeze_panes = "B2"
    wq.auto_filter.ref = f"A1:E{len(data['fin']) + 1}"

    # ---------- Hướng dẫn ----------
    n_rated = sum(x["rated"] for x in data["rows"])
    lines = [
        ("BẢNG ĐỊNH GIÁ CHUẨN VN100", F_TITLE),
        (
            (
                f"Ngày lập: {as_of:%d/%m/%Y} · {len(data['rows'])} mã · {n_rated} mã đủ điều kiện công bố giá trị hợp lý · "
                f"{len(data['reports'])} báo cáo CTCK trong {WINDOW_DAYS} ngày."
            ),
            F_BASE,
        ),
        ("", F_BASE),
        ("Cách đọc", F_BOLD),
        (
            (
                "• Tổng hợp: mỗi mã một dòng — giá trị hợp lý (ô vàng), upside, khuyến nghị, so sánh với CTCK, chỉ số cơ bản, "
                "chú ý và lưu ý định giá. Lọc cột 'Trạng thái công bố' = 'Công bố' để xem các mã có khuyến nghị."
            ),
            F_BASE,
        ),
        (
            "• Giá mục tiêu CTCK: từng báo cáo (CTCK, ngày, giá mục tiêu, khuyến nghị, tiêu đề, link) và upside tại ngày báo cáo.",
            F_BASE,
        ),
        (
            "• Giả định: rf, beta, COE, g, bội số mục tiêu, giả định ngân hàng (NIM, CIR, chi phí tín dụng, ROE bền vững).",
            F_BASE,
        ),
        (
            "• Số liệu TC: số liệu tài chính gốc (tỷ đồng) mà các chỉ số ở Tổng hợp tham chiếu.",
            F_BASE,
        ),
        (
            "• Phương pháp: cách tính từng phương pháp, tham số chung (rf, ERP, g, MOS) và quy tắc khuyến nghị.",
            F_BASE,
        ),
        (
            "• Chú thích cờ / Chất lượng dữ liệu: giải thích mã cờ và các mã dữ liệu còn thiếu/chưa sửa.",
            F_BASE,
        ),
        ("", F_BASE),
        ("Quy ước màu", F_BOLD),
        (
            (
                "Chữ xanh dương = số đầu vào/đầu ra engine (có thể sửa) · chữ đen = công thức · chữ xanh lá = tham chiếu sheet khác · "
                "nền vàng = giá trị hợp lý / tham số chính."
            ),
            F_BASE,
        ),
        ("", F_BASE),
        ("Lưu ý quan trọng", F_BOLD),
        (
            (
                "• 'Không xếp hạng' KHÔNG có nghĩa cổ phiếu xấu — mô hình chưa đủ cơ sở (thiếu dữ liệu dự án RNAV/SOTP, kết quả cực đoan "
                "chờ rà soát, dữ liệu lỗi). Không tự điền số cho các mã này."
            ),
            F_BASE,
        ),
        (
            "• Lệch lớn so với trung vị CTCK phải có giải thích ở cột 'Lưu ý định giá' — hệ thống không ép khớp với CTCK.",
            F_BASE,
        ),
        (
            (
                "• Ngân hàng: mô hình dùng chi phí tín dụng trung vị 3 năm và ràng buộc vốn E/A lịch sử — thận trọng hơn CTCK. "
                "Các mã bị ảnh hưởng mạnh có ghi chú 'Dự phóng LNST năm 1 ... thấp hơn TTM'."
            ),
            F_BASE,
        ),
        (
            "• Không có nhận định AI trong file này (khóa DeepSeek hiện không hợp lệ). Tiêu đề báo cáo CTCK chỉ để tham khảo nội bộ.",
            F_BASE,
        ),
        (
            "• Đây là công cụ phân tích nội bộ, không phải khuyến nghị đầu tư cá nhân hóa.",
            F_BASE,
        ),
        ("", F_BASE),
        (
            (
                "Nguồn: vn100_full.db (vnstock/VCI giá & BCTC; Simplize/24hMoney giá mục tiêu CTCK; TPCP 10 năm). "
                "Tái tạo: python -m scripts.export_valuation_book"
            ),
            F_BASE,
        ),
    ]
    for i, (text, font) in enumerate(lines, 1):
        _put(readme, i, 1, text, font=font)
    readme.column_dimensions["A"].width = 160

    from openpyxl.workbook.properties import CalcProperties

    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # Excel tự tính lại khi mở
    wb.save(path)


def main(out: str | None = None) -> None:
    as_of = _today_vn()
    path = (
        Path(out)
        if out
        else PROJECT_ROOT / "reports" / f"vn100_dinh_gia_chuan_{as_of:%Y%m%d}.xlsx"
    )
    db = SessionLocalRead()
    try:
        data = collect(db, as_of)
    finally:
        db.close()
    write_book(data, as_of, path)
    print(
        f"→ {path} | {len(data['rows'])} mã, {sum(x['rated'] for x in data['rows'])} công bố, "
        f"{len(data['reports'])} báo cáo CTCK"
    )


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
