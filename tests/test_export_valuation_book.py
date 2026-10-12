"""Bảng định giá chuẩn: trạng thái sửa BCTC và tham chiếu chéo giữa các sheet."""

import datetime
import os

import openpyxl

from scripts import export_valuation_book as book


def test_repair_status_db_evidence_beats_logs(tmp_path, monkeypatch):
    audit = tmp_path / "reports" / "audit"
    audit.mkdir(parents=True)
    old = audit / "bs_duplicate_repair_20260928.csv"
    new = audit / "bs_duplicate_repair_20260929_1048.csv"
    old.write_text(
        "ticker,action\nNT2,FETCH_ERROR ConnectionError\nHHV,FETCH_ERROR X\n",
        encoding="utf-8-sig",
    )
    new.write_text(
        "ticker,action\nNT2,INSERT\nTCH,FETCH_ERROR ConnectionError\n",
        encoding="utf-8-sig",
    )
    os.utime(old, (1, 1))
    monkeypatch.setattr(book, "PROJECT_ROOT", tmp_path)
    # HHV: log lỗi mạng nhưng DB đã có dòng tách (log lượt sửa thành công bị ghi đè).
    s = book._repair_status({"HHV"})
    assert s["HHV"].startswith("Đã sửa")
    assert "NT2" not in s  # lần sau chạy được, không còn lỗi
    assert s["TCH"].startswith("Chưa sửa")


def test_summary_price_column_matches_broker_sheet_reference(tmp_path):
    # 'Giá mục tiêu CTCK'!H tra giá hiện tại ở cột H của Tổng hợp; FV 60.000 / giá 50.000 → upside 20%.
    d = datetime.date(2026, 9, 29)
    row = {
        "ticker": "AAA",
        "name": "A",
        "exchange": "HOSE",
        "sector": "X",
        "nature": "Compounder",
        "method": "DCF/FCFF",
        "price_date": d,
        "price": 50_000.0,
        "intrinsic": 60_000.0,
        "relative": 60_000.0,
        "fv": 60_000.0,
        "mos": 0.15,
        "status": "Công bố",
        "cons_median": 55_000.0,
        "px_move": None,
        "caution": "",
        "note": "",
        "titles": "",
        "flags": "",
        "rated": True,
    }
    rep = {
        "ticker": "AAA",
        "broker": "B",
        "date": d,
        "tp": 55_000.0,
        "rating": "MUA",
        "px": 48_000.0,
        "used": "Có",
        "age": 0,
        "title": "",
        "site": "",
        "url": "",
    }
    data = {"rows": [row], "reports": [rep], "assumptions": [], "fin": [], "flags": {}}
    out = tmp_path / "b.xlsx"
    book.write_book(data, d, out)
    wb = openpyxl.load_workbook(out)
    ws, wr = wb["Tổng hợp"], wb["Giá mục tiêu CTCK"]
    assert ws["H1"].value == "Giá thị trường (VND)"
    assert "'Tổng hợp'!$H$2" in wr["H2"].value
    assert ws["L2"].value == '=IF(AND(ISNUMBER(K2),H2>0),K2/H2-1,"")'
    assert ws["K1"].value == "GIÁ TRỊ HỢP LÝ (VND)"
